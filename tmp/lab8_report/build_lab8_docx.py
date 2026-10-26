from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\Projects\ATSMA")
TMP = ROOT / "tmp" / "lab8_report"
TEMPLATE = Path(r"C:\College\ATSMA\Lab123BCB7137.docx")
OUTPUT = ROOT / "output" / "docx" / "Lab_8_Text_Clustering_KMeans_Report.docx"
FIGURES = TMP / "figures"


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def clear_body(document):
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def add_body(doc, text, size=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6):
    para = doc.add_paragraph()
    para.alignment = align
    para.paragraph_format.space_after = Pt(after)
    para.paragraph_format.line_spacing = 1.12
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    return para


def add_heading(doc, text, level=1):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    para.paragraph_format.space_after = Pt(5)
    run = para.add_run(text)
    run.bold = True
    run.font.size = Pt(14 if level == 1 else 11.5)
    return para


def add_bullet(doc, text):
    para = doc.add_paragraph(style="List Paragraph")
    para.paragraph_format.left_indent = Inches(0.25)
    para.paragraph_format.first_line_indent = Inches(-0.14)
    para.paragraph_format.space_after = Pt(3)
    run = para.add_run("• " + text)
    run.font.size = Pt(10.5)
    return para


def add_code(doc, text):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.15)
    para.paragraph_format.right_indent = Inches(0.1)
    para.paragraph_format.space_after = Pt(8)
    para.paragraph_format.line_spacing = 1.0
    run = para.add_run(text.strip())
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(8.3)
    return para


def add_figure(doc, filename, caption, width):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(FIGURES / filename), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    run = cap.add_run(caption)
    run.font.size = Pt(9)


def add_table(doc, rows, widths):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if c < len(row) - 1 else WD_ALIGN_PARAGRAPH.LEFT
            run = para.add_run(str(value))
            run.font.size = Pt(8.5)
            run.bold = r == 0
            if r == 0:
                run.font.color.rgb = RGBColor(255, 255, 255)
                shade(cell, "404040")
            elif r % 2 == 0:
                shade(cell, "F4F7FA")
            cell.width = Inches(widths[c])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    copy2(TEMPLATE, OUTPUT)
    doc = Document(OUTPUT)
    clear_body(doc)

    banner = doc.add_paragraph()
    banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
    banner.add_run().add_picture(str(ROOT / "tmp" / "atsma_report" / "template_banner.png"), width=Inches(6.5))
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    for label, value in [
        ("NAME:", " Vikhyat Shajee Nambiar"), ("REG NO:", " 23BCB7137"),
        ("COURSE CODE:", " CSE4039"), ("COURSE:", " Advanced Text and Social Media Analytics"),
        ("SLOT:", " L39-L40"), ("FACULTY NAME:", " Gundimeda Venugopal"),
    ]:
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(0)
        label_run = para.add_run(label)
        label_run.bold = True
        label_run.font.size = Pt(17)
        value_run = para.add_run(value)
        value_run.font.size = Pt(17)
    for title in ["LAB 8", "TEXT CLUSTERING USING K MEANS"]:
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(8)
        run = para.add_run(title)
        run.bold = True
        run.font.size = Pt(18)
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run("K-means text clustering report").font.size = Pt(10)
    doc.add_page_break()

    add_heading(doc, "Question:")
    add_body(doc, "Perform text document clustering with the K-means algorithm in Scikit-Learn. (A) Cluster news headlines and evaluate whether the discovered groups correspond to the sarcasm labels supplied in the dataset. (B) Cluster national anthems to examine recurring lexical themes such as nationhood, freedom, land, religion, and patriotic imagery.")
    add_heading(doc, "Dataset and implementation details:")
    add_bullet(doc, "Sarcasm dataset: 26,709 English news headlines, each with an is_sarcastic label used only for post-clustering evaluation.")
    add_bullet(doc, "National anthems dataset: 190 anthem texts with country and continent metadata.")
    add_bullet(doc, "Text representation: TF-IDF vectors with English stop words removed and unigram/bigram features.")
    add_bullet(doc, "Clustering: KMeans with random_state=42 and multiple initialisations for reproducibility.")
    add_bullet(doc, "Evaluation: cluster sizes, top centroid terms, cosine silhouette score, and an external label comparison for the sarcasm dataset.")
    add_heading(doc, "Method:")
    add_body(doc, "K-means is an unsupervised method. It does not learn the sarcasm label; it partitions TF-IDF vectors according to lexical similarity. Therefore, the sarcasm label is used only after clustering to measure how closely lexical clusters agree with the provided annotation. This distinction is important because sarcasm often depends on context and pragmatic contrast rather than words alone.")

    doc.add_page_break()
    add_heading(doc, "A) Sarcasm headline clustering", 1)
    add_heading(doc, "Code:")
    add_code(doc, '''import json
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.cluster import KMeans
from sklearn.metrics import silhouette_score

records = json.load(open("sarcasm.json", encoding="utf-8"))
headlines = [row["headline"] for row in records]
labels = [row["is_sarcastic"] for row in records]

vectorizer = TfidfVectorizer(stop_words="english", ngram_range=(1, 2),
                             min_df=3, max_features=5000)
X = vectorizer.fit_transform(headlines)
kmeans = KMeans(n_clusters=2, n_init=20, random_state=42)
clusters = kmeans.fit_predict(X)
silhouette = silhouette_score(X[:2500], clusters[:2500], metric="cosine")''')
    add_heading(doc, "Working of the code:")
    add_bullet(doc, "Headlines are converted to 5,000 TF-IDF features. Common English stop words are removed; bigrams retain phrases such as 'donald trump'.")
    add_bullet(doc, "K=2 is selected because the supplied dataset has two labels: sarcastic and not sarcastic.")
    add_bullet(doc, "The two cluster labels are compared with the dataset labels through a contingency table; K-means label numbers themselves have no semantic meaning.")
    add_heading(doc, "Output:")
    add_table(doc, [
        ["Cluster", "Headlines", "Not sarcastic", "Sarcastic", "Top centroid terms"],
        ["0", "23,602", "12,932", "10,670", "man, area, just, woman, report, year"],
        ["1", "3,107", "2,053", "1,054", "trump, new, donald, new york, introduces"],
    ], [0.6, 0.8, 1.0, 0.8, 3.25])
    add_body(doc, "The cosine silhouette score is 0.0038, which indicates very weak separation between the two lexical clusters. Mapping each cluster to its majority dataset label yields 56.1% agreement, and both clusters are majority non-sarcastic. The result demonstrates that headline vocabulary alone is not a reliable unsupervised detector of sarcasm.")
    add_figure(doc, "sarcasm_clusters.png", "Figure 1. K-means clusters compared with the supplied sarcasm labels.", 6.0)
    add_heading(doc, "Interpretation:")
    add_body(doc, "The smaller cluster is strongly associated with political vocabulary, especially terms around Donald Trump, rather than a distinct sarcastic class. This is a meaningful diagnostic result: K-means has grouped headlines by subject matter, while sarcasm requires contextual and semantic signals that are not captured well by bag-of-words TF-IDF features.")

    doc.add_page_break()
    add_heading(doc, "B) National anthem clustering", 1)
    add_heading(doc, "Code:")
    add_code(doc, '''import pandas as pd
from sklearn.decomposition import TruncatedSVD

anthems = pd.read_csv("anthems.csv").dropna(subset=["Anthem"])
vectorizer = TfidfVectorizer(stop_words="english", ngram_range=(1, 2),
                             min_df=2, max_features=3000)
X = vectorizer.fit_transform(anthems["Anthem"])
kmeans = KMeans(n_clusters=5, n_init=50, random_state=42)
anthems["cluster"] = kmeans.fit_predict(X)

svd = TruncatedSVD(n_components=2, random_state=42)
coordinates = svd.fit_transform(X)''')
    add_heading(doc, "Working of the code:")
    add_bullet(doc, "The anthem texts are vectorised with 1,830 retained TF-IDF features after preprocessing.")
    add_bullet(doc, "K=5 creates an interpretable set of broad lexical groups without imposing country or continent labels on the model.")
    add_bullet(doc, "The highest-weight centroid terms provide a short label for each cluster, while TruncatedSVD projects the sparse document matrix to two dimensions for visual inspection.")
    add_heading(doc, "Output:")
    add_table(doc, [
        ["Cluster", "Anthems", "Top terms", "Example countries"],
        ["0", "24", "thy, thee, land, god, love", "Azerbaijan, Cyprus, France, Greece"],
        ["1", "34", "country, flag, people, fatherland, glory", "Albania, Armenia, Germany, Malta"],
        ["2", "35", "homeland, thou, hail, beloved, sea", "Bulgaria, Croatia, Denmark, Hungary"],
        ["3", "39", "land, free, home, god, bless", "Austria, Belarus, Bosnia and Herzegovina"],
        ["4", "58", "freedom, people, africa, nation, liberty", "Belgium, Georgia, Ireland, Italy"],
    ], [0.55, 0.6, 2.65, 2.65])
    add_body(doc, "The five-cluster cosine silhouette score is 0.0127. The low score shows considerable lexical overlap, which is expected because national anthems share ceremonial language. Nevertheless, centroid terms expose recurring themes of homeland, faith, freedom, people, and national symbols.")
    add_figure(doc, "anthem_clusters.png", "Figure 2. National anthems in a two-dimensional latent-semantic projection, coloured by K-means cluster.", 5.7)

    doc.add_page_break()
    add_heading(doc, "Conclusion:")
    add_body(doc, "K-means with TF-IDF provides a transparent first approach to grouping documents by lexical content. For sarcasm headlines, the clusters mainly reflect topic vocabulary and do not align strongly with sarcasm labels, so a supervised contextual model would be more appropriate for detection. For national anthems, the clusters expose broad recurring vocabulary about land, nationhood, freedom, faith, and patriotic identity, although the low silhouette score confirms that these themes overlap across countries.")
    add_heading(doc, "Limitations and future work:")
    add_bullet(doc, "K-means requires a value for K; different choices can produce different thematic granularity.")
    add_bullet(doc, "TF-IDF does not capture word order, irony, or cultural context; sentence embeddings could improve the representation.")
    add_bullet(doc, "Sarcasm detection should be evaluated with a supervised classifier and held-out test set, rather than relying on unsupervised clusters.")
    add_heading(doc, "References:")
    add_body(doc, "GeeksforGeeks. Clustering Text Documents Using K-Means in Scikit-Learn. https://www.geeksforgeeks.org/machine-learning/clustering-text-documents-using-k-means-in-scikit-learn/")
    add_body(doc, "PawanKrGunjan. Sarcasm Detection dataset. https://raw.githubusercontent.com/PawanKrGunjan/Natural-Language-Processing/main/Sarcasm%20Detection/sarcasm.json")
    add_body(doc, "Lucas Turtle. National Anthems of the World dataset. https://www.kaggle.com/lucasturtle/national-anthems-of-the-world")
    add_body(doc, "Lucas De Sa. Text Clustering with K-Means. https://medium.com/@lucasdesa/text-clustering-with-k-means-a039d84a941b")
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
